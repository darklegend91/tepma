"""Render a profile dict into an editable Word (.docx) resume."""
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ACCENT = RGBColor(0x28, 0x3C, 0x6E)
MUTED = RGBColor(0x6E, 0x6E, 0x6E)


def _section(doc: Document, title: str): # type:ignore
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = ACCENT


def _entry(doc: Document, left: str, right: str = ""): # type:ignore
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(left)
    run.bold = True
    run.font.size = Pt(11)
    if right:
        tail = p.add_run(f"   ({right})")
        tail.font.size = Pt(10)
        tail.font.color.rgb = MUTED


def render_resume_docx(profile: dict) -> bytes:
    doc = Document()
    for section in doc.sections:
        for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(section, attr, Pt(50))

    # Header
    p = doc.add_paragraph()
    run = p.add_run(profile.get("name", ""))
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = ACCENT
    if profile.get("target_role"):
        p = doc.add_paragraph()
        run = p.add_run(profile["target_role"])
        run.font.size = Pt(13)
        run.font.color.rgb = MUTED
    contact = "  |  ".join(
        v for v in (profile.get("email"), profile.get("phone"), profile.get("location")) if v
    )
    if contact:
        p = doc.add_paragraph()
        run = p.add_run(contact)
        run.font.size = Pt(10)
        run.font.color.rgb = MUTED

    if profile.get("summary"):
        _section(doc, "Summary")
        doc.add_paragraph(profile["summary"])

    if profile.get("experience"):
        _section(doc, "Experience")
        for job in profile["experience"]:
            dates = " - ".join(v for v in (job.get("start"), job.get("end")) if v)
            _entry(doc, f"{job.get('title', '')} — {job.get('company', '')}", dates)
            for b in job.get("bullets", []):
                doc.add_paragraph(b, style="List Bullet")

    if profile.get("projects"):
        _section(doc, "Projects")
        for proj in profile["projects"]:
            _entry(doc, proj.get("name", ""), proj.get("technologies", ""))
            doc.add_paragraph(proj.get("description", ""))

    if profile.get("education"):
        _section(doc, "Education")
        for edu in profile["education"]:
            _entry(doc, f"{edu.get('degree', '')} — {edu.get('institution', '')}", edu.get("year", ""))
            if edu.get("details"):
                doc.add_paragraph(edu["details"])

    if profile.get("skills"):
        _section(doc, "Skills")
        doc.add_paragraph(", ".join(profile["skills"]))

    if profile.get("achievements"):
        _section(doc, "Achievements")
        for a in profile["achievements"]:
            doc.add_paragraph(a, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
