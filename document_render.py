"""Render a user-authored application/letter as printable PDF and editable DOCX."""
import io
import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fpdf import FPDF


MARGIN = 22
BODY = (32, 37, 48)
MUTED = (100, 110, 125)


def _font_path() -> Path | None:
    """Find a Unicode font without making deployment depend on one operating system."""
    configured = os.getenv("TEPMA_UNICODE_FONT", "").strip()
    candidates = [
        Path(configured) if configured else None,
        Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf"),
    ]
    return next((path for path in candidates if path is not None and path.is_file()), None)


def _latin(text: str) -> str:
    return (text or "").encode("latin-1", "replace").decode("latin-1")


def _lines(value) -> list[str]:
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    return [str(value).strip()] if str(value or "").strip() else []


def render_application_pdf(document: dict) -> bytes:
    """Create an A4 PDF from the constrained document JSON returned by the LLM."""
    pdf = FPDF("P", "mm", "A4")
    pdf.set_margins(MARGIN, MARGIN, MARGIN)
    pdf.set_auto_page_break(True, margin=MARGIN)
    pdf.add_page()

    unicode_font = _font_path()
    if unicode_font:
        pdf.add_font("TePMAUnicode", fname=str(unicode_font))
        family = "TePMAUnicode"
        clean = lambda value: str(value or "")
    else:
        family = "helvetica"
        clean = lambda value: _latin(str(value or ""))

    pdf.set_text_color(*BODY)
    pdf.set_font(family, size=16)
    pdf.multi_cell(
        0, 8, clean(document.get("title", "Application")), align="C",
        new_x="LMARGIN", new_y="NEXT",
    )
    pdf.ln(5)

    date_text = clean(document.get("date", ""))
    if date_text:
        pdf.set_font(family, size=10)
        pdf.set_text_color(*MUTED)
        pdf.cell(0, 6, date_text, align="R", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)

    pdf.set_text_color(*BODY)
    pdf.set_font(family, size=11)
    for line in _lines(document.get("recipient_lines", [])):
        pdf.multi_cell(0, 6, clean(line), new_x="LMARGIN", new_y="NEXT")

    subject = str(document.get("subject", "")).strip()
    if subject:
        pdf.ln(3)
        pdf.set_font(family, size=11)
        pdf.multi_cell(
            0, 6, clean(f"Subject: {subject}"), new_x="LMARGIN", new_y="NEXT"
        )

    pdf.ln(4)
    salutation = str(document.get("salutation", "Sir/Madam,")).strip()
    if salutation:
        pdf.multi_cell(0, 6, clean(salutation), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)

    for paragraph in _lines(document.get("paragraphs", [])):
        pdf.multi_cell(
            0, 6, clean(paragraph), align="J", new_x="LMARGIN", new_y="NEXT"
        )
        pdf.ln(3)

    closing = str(document.get("closing", "Yours sincerely,")).strip()
    if closing:
        pdf.ln(2)
        pdf.multi_cell(0, 6, clean(closing), new_x="LMARGIN", new_y="NEXT")
    for line in _lines(document.get("sender_lines", [])):
        pdf.multi_cell(0, 6, clean(line), new_x="LMARGIN", new_y="NEXT")

    return bytes(pdf.output())


def render_application_docx(document: dict) -> bytes:
    """Create an editable Word copy of the generated application."""
    doc = Document()
    for section in doc.sections:
        for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(section, attr, Pt(62))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(str(document.get("title", "Application")))
    run.bold = True
    run.font.size = Pt(16)

    date_text = str(document.get("date", "")).strip()
    if date_text:
        paragraph = doc.add_paragraph(date_text)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for line in _lines(document.get("recipient_lines", [])):
        doc.add_paragraph(line)

    subject = str(document.get("subject", "")).strip()
    if subject:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(f"Subject: {subject}")
        run.bold = True

    salutation = str(document.get("salutation", "Sir/Madam,")).strip()
    if salutation:
        doc.add_paragraph(salutation)
    for body in _lines(document.get("paragraphs", [])):
        paragraph = doc.add_paragraph(body)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    closing = str(document.get("closing", "Yours sincerely,")).strip()
    if closing:
        doc.add_paragraph(closing)
    for line in _lines(document.get("sender_lines", [])):
        doc.add_paragraph(line)

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
